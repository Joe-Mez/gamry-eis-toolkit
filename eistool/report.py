"""Excel workbooks, CSV and Word tables."""
from __future__ import annotations

import math
import re
from pathlib import Path

import numpy as np
import pandas as pd

from .circuits import ELEMENTS
from .dta import _num


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\[\]\:\*\?\/\\]", "_", name)[:31] or "Sheet"
    s, k = base, 2
    while s.lower() in used:
        suffix = f"_{k}"
        s = base[: 31 - len(suffix)] + suffix
        k += 1
    used.add(s.lower())
    return s


def system_dataframe(sysrec) -> pd.DataFrame:
    """One table with measured data (+ fit and residuals when available)."""
    d = sysrec.data
    df = pd.DataFrame({
        "Frequency (Hz)": d.freq,
        "Z' (Ω)": d.zreal,
        "Z'' (Ω)": d.zimag,
        "-Z'' (Ω)": -d.zimag,
        "|Z| (Ω)": d.zmod,
        "Phase (°)": d.phase_deg,
    })
    if sysrec.area_used:
        a = sysrec.area_used
        df["Z' (Ω cm²)"] = d.zreal * a
        df["-Z'' (Ω cm²)"] = -d.zimag * a
        df["|Z| (Ω cm²)"] = d.zmod * a
    if sysrec.fit is not None:
        u = " (Ω cm²)" if sysrec.area_used else " (Ω)"
        zf = sysrec.fit.predict(d.freq)
        zmeas = sysrec.z_fit_units
        df["Fit Z'" + u] = zf.real
        df["Fit -Z''" + u] = -zf.imag
        df["Fit |Z|" + u] = np.abs(zf)
        df["Fit -Phase (°)"] = -np.degrees(np.angle(zf))
        in_range = (d.freq >= sysrec.fit.freq_range[0]) & (d.freq <= sysrec.fit.freq_range[1])
        with np.errstate(invalid="ignore", divide="ignore"):
            df["Residual Z' (%)"] = np.where(in_range, 100 * (zmeas.real - zf.real) / np.abs(zmeas), np.nan)
            df["Residual Z'' (%)"] = np.where(in_range, 100 * (zmeas.imag - zf.imag) / np.abs(zmeas), np.nan)
    return df


def fit_table(systems, param_labels: dict | None = None) -> pd.DataFrame:
    """Wide table: one row per system, value + error columns for each parameter."""
    param_labels = param_labels or {}
    order: list[tuple[str, str]] = []  # (param name, unit)
    for s in systems:
        if s.fit is None:
            continue
        for n, u in zip(s.fit.circuit.param_names, s.fit.circuit.units(bool(s.area_used))):
            if n not in [o[0] for o in order]:
                order.append((n, u))
    rows = []
    for s in systems:
        if s.fit is None:
            continue
        r = {"System": s.name, "Legend label": s.label, "Circuit": s.fit.circuit.text}
        p = s.fit.params
        err = dict(zip(s.fit.circuit.param_names, s.fit.rel_error_pct))
        for n, u in order:
            lab = param_labels.get(n, n)
            col = f"{lab} ({u})" if u not in ("–", "") else lab
            r[col] = p.get(n, np.nan)
            r[f"{lab} error (%)"] = err.get(n, np.nan)
        r["χ²"] = s.fit.chi2
        r["Weighting"] = s.fit.weighting
        r["Points"] = s.fit.n_points
        r["f min (Hz)"] = s.fit.freq_range[0]
        r["f max (Hz)"] = s.fit.freq_range[1]
        r["Area used (cm²)"] = s.area_used or ""
        rows.append(r)
    return pd.DataFrame(rows)


def _autofit(ws, df: pd.DataFrame):
    from openpyxl.styles import Alignment, Font
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(df.columns, start=1):
        width = max(10, min(40, len(str(col)) + 2))
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = "A2"


def write_workbook(path: Path, systems, fits_table: pd.DataFrame | None = None):
    path.parent.mkdir(parents=True, exist_ok=True)
    used: set = set()
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        if fits_table is not None and len(fits_table):
            name = _sheet_name("Fit results", used)
            fits_table.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            _autofit(ws, fits_table)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "0.000E+00"

        # All systems side by side (handy for Origin / Excel plotting)
        blocks = []
        for s in systems:
            u = "Ω cm²" if s.area_used else "Ω"
            z = s.z_fit_units
            blocks.append(pd.DataFrame({
                (s.name, "Frequency (Hz)"): s.data.freq,
                (s.name, f"Z' ({u})"): z.real,
                (s.name, f"-Z'' ({u})"): -z.imag,
                (s.name, f"|Z| ({u})"): np.abs(z),
                (s.name, "-Phase (°)"): -np.degrees(np.angle(z)),
            }))
        allsys = pd.concat(blocks, axis=1)
        flat = allsys.copy()
        flat.columns = [f"{a} | {b}" for a, b in allsys.columns]
        name = _sheet_name("All systems", used)
        flat.to_excel(xw, sheet_name=name, index=False)
        _autofit(xw.sheets[name], flat)

        for s in systems:
            df = system_dataframe(s)
            name = _sheet_name(s.name, used)
            df.to_excel(xw, sheet_name=name, index=False)
            _autofit(xw.sheets[name], df)

        info = pd.DataFrame([{
            "System": s.name, "Legend label": s.label, "Source file": s.data.path.name,
            "Area in DTA (cm²)": s.data.area, "Area used (cm²)": s.area_used or "none (raw Ω)",
            "Points": len(s.data.freq), "f max (Hz)": s.data.freq.max(), "f min (Hz)": s.data.freq.min(),
            "Gamry TITLE": s.data.header.get("TITLE", ""), "Date": s.data.header.get("DATE", ""),
            "Eoc (V)": _num(s.data.header.get("EOC", "")) if s.data.header.get("EOC") else "",
        } for s in systems])
        name = _sheet_name("Info", used)
        info.to_excel(xw, sheet_name=name, index=False)
        _autofit(xw.sheets[name], info)
    return path


# ---------------------------------------------------------------------------
# Word table (manuscript ready)
# ---------------------------------------------------------------------------

_SUP = str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")


def fmt_value(v: float, sig: int = 3) -> tuple[str, str | None]:
    """Return (mantissa text, exponent or None)."""
    if v is None or not np.isfinite(v):
        return "–", None
    if v == 0:
        return "0", None
    e = int(math.floor(math.log10(abs(v))))
    if -2 <= e < 4:
        decimals = max(sig - 1 - e, 0)
        return f"{v:.{decimals}f}", None
    m = v / 10 ** e
    if round(abs(m), sig - 1) >= 10:
        m, e = m / 10, e + 1
    return f"{m:.{sig - 1}f}", str(e)


def fmt_value_text(v: float, sig: int = 3) -> str:
    m, e = fmt_value(v, sig)
    return m if e is None else f"{m} × 10{e.translate(_SUP)}"


def _add_rich(paragraph, text: str):
    """Add text where 'X_sub' becomes subscript and '^sup' superscript (single token)."""
    for part in re.split(r"(_\{[^}]*\}|_\w+|\^\{[^}]*\}|\^[\w\-⁻]+)", text):
        if not part:
            continue
        if part.startswith("_"):
            r = paragraph.add_run(part[1:].strip("{}"))
            r.font.subscript = True
        elif part.startswith("^"):
            r = paragraph.add_run(part[1:].strip("{}"))
            r.font.superscript = True
        else:
            paragraph.add_run(part)


def write_docx_table(path: Path, systems, param_labels: dict | None, show_errors: bool,
                     caption: str | None = None, font: str = "Arial"):
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
    except ImportError:
        return None

    param_labels = param_labels or {}
    fitted = [s for s in systems if s.fit is not None]
    if not fitted:
        return None
    cols: list[tuple[str, str]] = []
    for s in fitted:
        for n, u in zip(s.fit.circuit.param_names, s.fit.circuit.units(bool(s.area_used))):
            if n not in [c[0] for c in cols]:
                cols.append((n, u))

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(9)

    circuits = sorted({s.fit.circuit.text for s in fitted})
    cap = caption or ("Electrochemical parameters obtained by fitting the EIS data with the "
                      "equivalent circuit " + ", ".join(circuits) + ".")
    p = doc.add_paragraph()
    r = p.add_run("Table 1. ")
    r.bold = True
    p.add_run(cap)

    headers = ["Sample"] + [c[0] for c in cols] + ["χ²"]
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (h, cell) in enumerate(zip(headers, t.rows[0].cells)):
        para = cell.paragraphs[0]
        if 0 < i < len(headers) - 1:
            n, u = cols[i - 1]
            _add_rich(para, param_labels.get(n, _default_label(n)))
            if u not in ("–", ""):
                para.add_run("\n(" + u + ")")
        else:
            para.add_run(h)
        for run in para.runs:
            run.bold = True

    for s in fitted:
        cells = t.add_row().cells
        _add_rich(cells[0].paragraphs[0], s.label)
        params = s.fit.params
        err = dict(zip(s.fit.circuit.param_names, s.fit.rel_error_pct))
        for j, (n, _) in enumerate(cols, start=1):
            para = cells[j].paragraphs[0]
            m, e = fmt_value(params.get(n, np.nan))
            para.add_run(m)
            if e is not None:
                para.add_run(" × 10")
                sup = para.add_run(e.replace("-", "−"))
                sup.font.superscript = True
            if show_errors and n in err and np.isfinite(err[n]):
                ev = err[n]
                para.add_run(" (<0.1)" if ev < 0.1 else f" ({ev:.1f})" if ev < 100 else f" ({ev:.0f})")
        m, e = fmt_value(s.fit.chi2)
        para = cells[-1].paragraphs[0]
        para.add_run(m)
        if e is not None:
            para.add_run(" × 10")
            sup = para.add_run(e.replace("-", "−"))
            sup.font.superscript = True

    # three-line table (journal style): top rule, header rule, bottom rule
    def border(cell, **edges):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for edge, val in edges.items():
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(val))
            el.set(qn("w:color"), "000000")
            b.append(el)
        tcPr.append(b)

    for cell in t.rows[0].cells:
        border(cell, top=8, bottom=4)
    for cell in t.rows[-1].cells:
        border(cell, bottom=8)
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(8)
                    run.font.name = font

    if show_errors:
        doc.add_paragraph("Values in parentheses are relative standard errors of the fitted "
                          "parameters (%). χ² is the weighted sum of squares divided by the "
                          "degrees of freedom (" + fitted[0].fit.weighting + " weighting).")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def _default_label(name: str) -> str:
    """R1 -> R_1, CPE1_Q -> Q_1, CPE1_n -> n_1"""
    m = re.match(r"(CPE|Ws|Wo|R|C|L|W)(\w*?)(?:_(Q|n|R|T))?$", name)
    if not m:
        return name
    et, idx, suf = m.groups()
    if et == "CPE":
        base = suf
    elif et in ("Ws", "Wo"):
        base = f"{et}-{suf}" if suf else et
    else:
        base = et
    return f"{base}_{{{idx}}}" if idx else base


def element_help() -> str:
    return "\n".join(f"{k}: {v['describe']}" for k, v in ELEMENTS.items())
